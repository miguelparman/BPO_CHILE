"""EXTRACCION: equivalente a cada 'Origen de Excel' del .dtsx
(Microsoft.ACE.OLEDB, 'EXCEL 12.0 XML;HDR=YES', OpenRowset 'Hoja1$').

Reproduce como el driver ACE entrega los datos al Data Flow:
  - HDR=YES: la fila 1 es la cabecera; se usan las 8 columnas que mapea el
    paquete (si falta alguna, SSIS falla en validacion -> ExtraccionError).
  - Cada columna tiene UN tipo (el del metadata del .dtsx): 'Caso SF'
    DT_R8, 'Fecha' DT_DATE, el resto DT_WSTR(255). Sin IMEX, ACE devuelve
    NULL para las celdas de otro tipo ("tipo minoritario"): un 'SIN CASO'
    en 'Caso SF', un texto en 'Fecha' o un numero en una columna de texto
    llegan como NULL. Las filas con 'Caso SF' NULL las elimina despues el
    DELETE del paquete.
  - Se leen los valores calculados de las formulas (data_only), como ACE.
  - Se devuelven todas las filas del rango usado, incluidas las vacias
    (llegan con todo NULL y tambien las borra el DELETE).

Fuentes: FuenteSharePoint (normal: ReportingFractalia via Graph) y
FuenteLocal (contingencia: una carpeta en disco con las 4 subcarpetas).
"""

from __future__ import annotations

import io
import logging
from datetime import date, datetime
from pathlib import Path
from typing import Any, Protocol

import openpyxl

from exceptions import ExtraccionError
from logging_setup import NOMBRE_LOGGER
from mappings import CASO, FECHA, HOJA, TEXTO, FlujoSpec
from sharepoint.client import SharePointClient

logger = logging.getLogger(NOMBRE_LOGGER)

Fila = dict[str, Any]  # {columna origen: valor tipado como lo entrega ACE}


class FuenteArchivo(Protocol):
    def leer(self, flujo: FlujoSpec) -> bytes: ...

    def describir(self, flujo: FlujoSpec) -> str: ...


class ProveedorGraph(Protocol):
    def obtener(self) -> tuple[SharePointClient, str]: ...


class FuenteSharePoint:
    """Lee el Excel ya copiado en 'ReportingFractalia' (origen oficial)."""

    def __init__(self, destino: ProveedorGraph) -> None:
        self._destino = destino

    def leer(self, flujo: FlujoSpec) -> bytes:
        cliente, drive_id = self._destino.obtener()
        return cliente.download_file(drive_id, flujo.ruta_reporting)

    def describir(self, flujo: FlujoSpec) -> str:
        return f"SharePoint ReportingFractalia: {flujo.ruta_reporting}"


class FuenteLocal:
    """Contingencia: <raiz>/<subcarpeta>/<archivo> (misma estructura que
    '...\\BPO - Insumos\\Chile\\Especializados', la ruta original del .dtsx)."""

    def __init__(self, raiz: Path) -> None:
        self._raiz = raiz

    def _ruta(self, flujo: FlujoSpec) -> Path:
        return self._raiz / flujo.subcarpeta / flujo.archivo

    def leer(self, flujo: FlujoSpec) -> bytes:
        ruta = self._ruta(flujo)
        if not ruta.exists():
            raise ExtraccionError(f"No existe el archivo local: {ruta}")
        return ruta.read_bytes()

    def describir(self, flujo: FlujoSpec) -> str:
        return f"archivo local: {self._ruta(flujo)}"


def valor_ace(tipo: str, valor: Any) -> Any:
    """Valor de una celda tal como lo entrega ACE para una columna del tipo
    dado; NULL (None) si la celda es de otro tipo."""
    if valor is None:
        return None
    if tipo == CASO:  # DT_R8
        if isinstance(valor, bool) or not isinstance(valor, (int, float)):
            return None
        return float(valor)
    if tipo == FECHA:  # DT_DATE
        if isinstance(valor, datetime):
            return valor
        if isinstance(valor, date):
            return datetime(valor.year, valor.month, valor.day)
        return None
    if tipo == TEXTO:  # DT_WSTR(255)
        return valor if isinstance(valor, str) else None
    raise ValueError(f"Tipo de columna desconocido: {tipo}")


def leer_hoja(contenido: bytes, flujo: FlujoSpec) -> list[Fila]:
    """Lee 'Hoja1' y devuelve una fila por cada fila de datos (bajo la
    cabecera) con las 8 columnas del flujo tipadas como ACE."""
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # zip corrupto, no es xlsx, etc.
        raise ExtraccionError(f"[{flujo.nombre}] No se pudo abrir '{flujo.archivo}': {exc}") from exc
    try:
        if HOJA not in libro.sheetnames:
            raise ExtraccionError(f"[{flujo.nombre}] '{flujo.archivo}' no tiene la hoja '{HOJA}' (hojas: {libro.sheetnames}).")
        filas_excel = libro[HOJA].iter_rows(values_only=True)
        cabecera = next(filas_excel, None) or ()
        indices: dict[str, int] = {}
        for i, nombre in enumerate(cabecera):
            if isinstance(nombre, str) and nombre not in indices:
                indices[nombre] = i  # ante nombres repetidos ACE usa la primera
        faltantes = [c for c in flujo.columnas_origen if c not in indices]
        if faltantes:
            raise ExtraccionError(
                f"[{flujo.nombre}] '{flujo.archivo}' / {HOJA}: faltan las columnas {faltantes} (cabecera: {cabecera})."
            )
        filas: list[Fila] = []
        for fila in filas_excel:
            filas.append(
                {
                    c.origen: valor_ace(c.tipo, fila[indices[c.origen]] if indices[c.origen] < len(fila) else None)
                    for c in flujo.columnas
                }
            )
        return filas
    finally:
        libro.close()


class Extractor:
    def __init__(self, fuente: FuenteArchivo) -> None:
        self._fuente = fuente

    def extraer(self, flujo: FlujoSpec) -> list[Fila]:
        logger.info("[%s] Leyendo %s", flujo.nombre, self._fuente.describir(flujo))
        filas = leer_hoja(self._fuente.leer(flujo), flujo)
        logger.info("[%s] %s fila(s) leidas de '%s'.", flujo.nombre, len(filas), HOJA)
        return filas
