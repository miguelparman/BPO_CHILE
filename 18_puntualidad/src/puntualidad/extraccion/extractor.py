r"""Extraccion: el Excel SLA (Excel Source 'SLA' del Data Flow TBL_SLA_CHILE).

En SSIS: Connection Manager Excel 'SLA' (ACE.OLEDB.12.0, EXCEL 12.0 XML, HDR=YES)
sobre ...\BPO - Insumos\Chile\INDICADORES CHILE\SLA Puntualidad\DB.xlsx, hoja
'SLA - Carterizado$'. Aqui el mismo DB.xlsx se descarga de SharePoint.

  - FuenteSharePoint / FuenteLocal entregan los bytes del .xlsx.
  - leer_excel() es el Excel Source: encabezados + filas de la hoja, tal cual
    (sin tipar; la conversion de tipos es transformacion/transformer.py).
"""

from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Protocol

from openpyxl import load_workbook

from exceptions import ExtraccionError
from logging_setup import NOMBRE_LOGGER

logger = logging.getLogger(NOMBRE_LOGGER)


class FuenteExcel(Protocol):
    def leer(self) -> bytes: ...

    def describir(self) -> str: ...


class FuenteSharePoint:
    """DB.xlsx en ReportingFractalia (carpeta '08 PUNTUALIDAD')."""

    def __init__(self, cliente, drive_id: str, ruta: str, web_drive: str = "") -> None:
        self.cliente, self.drive_id, self.ruta = cliente, drive_id, ruta
        self._web_drive = web_drive

    def leer(self) -> bytes:
        try:
            return self.cliente.download_file(self.drive_id, self.ruta)
        except Exception as exc:
            raise ExtraccionError(f"No se pudo descargar '{self.ruta}' de SharePoint: {exc}") from exc

    def describir(self) -> str:
        return f"{self._web_drive}/{self.ruta}"


class FuenteLocal:
    """Contingencia / pruebas: el .xlsx en una ruta local."""

    def __init__(self, ruta: Path) -> None:
        self._ruta = Path(ruta)

    def leer(self) -> bytes:
        try:
            return self._ruta.read_bytes()
        except OSError as exc:
            raise ExtraccionError(f"No se pudo leer '{self._ruta}': {exc}") from exc

    def describir(self) -> str:
        return str(self._ruta)


def leer_excel(contenido: bytes, hoja: str) -> tuple[list[str], list[tuple]]:
    """(encabezados, filas) de la hoja, con los valores de celda tal cual
    (texto, numero, fecha o None). La primera fila es el encabezado
    (HDR=YES). Las filas totalmente vacias se omiten con una advertencia."""
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # archivo corrupto, no es xlsx, etc.
        raise ExtraccionError(f"No se pudo abrir el Excel: {exc}") from exc
    try:
        if hoja not in libro.sheetnames:
            raise ExtraccionError(f"El Excel no tiene la hoja '{hoja}'. Hojas: {libro.sheetnames}")
        filas = list(libro[hoja].iter_rows(values_only=True))
    finally:
        libro.close()
    if not filas:
        raise ExtraccionError(f"La hoja '{hoja}' esta vacia (sin fila de encabezado).")
    encabezados = ["" if c is None else str(c) for c in filas[0]]
    datos = [tuple(f) for f in filas[1:]]
    vacias = sum(1 for f in datos if all(v is None for v in f))
    if vacias:
        logger.warning("Hoja '%s': se omiten %s fila(s) totalmente vacias.", hoja, vacias)
        datos = [f for f in datos if not all(v is None for v in f)]
    return encabezados, datos
