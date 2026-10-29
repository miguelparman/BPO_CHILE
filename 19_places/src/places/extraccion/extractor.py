"""EXTRACCION: equivalente al 'Origen de Excel' del .dtsx
(Microsoft.ACE.OLEDB, 'EXCEL 12.0 XML;HDR=YES', OpenRowset 'Hoja1$').

Reproduce como el driver ACE entrega los datos al Data Flow:
  - HDR=YES: la fila 1 es la cabecera; se usan las 7 columnas del metadata
    del .dtsx (si falta alguna, SSIS falla en validacion -> ExtraccionError).
  - Las 7 columnas son DT_WSTR(255). Sin IMEX, ACE devuelve NULL para las
    celdas de otro tipo ("tipo minoritario"): un numero en una columna de
    texto llega como NULL. Las filas con 'id' NULL las elimina despues
    'LIMPIA TABLA'.
  - Se leen los valores calculados de las formulas (data_only), como ACE.
  - Se devuelven todas las filas del rango usado, incluidas las vacias
    (llegan con todo NULL y tambien las borra 'LIMPIA TABLA').

Fuentes: FuenteSharePoint (normal: ReportingFractalia via Graph) y
FuenteLocal (contingencia: un .xlsx en disco, la ruta original del .dtsx).
"""

from __future__ import annotations

import io
import logging
from pathlib import Path
from typing import Any, Protocol

import openpyxl

import mappings
from exceptions import ExtraccionError
from logging_setup import NOMBRE_LOGGER
from sharepoint.client import SharePointClient

logger = logging.getLogger(NOMBRE_LOGGER)

Fila = dict[str, Any]  # {columna origen: valor tipado como lo entrega ACE}


class FuenteArchivo(Protocol):
    def leer(self) -> bytes: ...

    def describir(self) -> str: ...


class ProveedorGraph(Protocol):
    def obtener(self) -> tuple[SharePointClient, str]: ...


class FuenteSharePoint:
    """Lee el Excel desde 'ReportingFractalia' (reemplaza la ruta en disco del .dtsx)."""

    def __init__(self, sitio: ProveedorGraph) -> None:
        self._sitio = sitio

    def leer(self) -> bytes:
        cliente, drive_id = self._sitio.obtener()
        return cliente.download_file(drive_id, mappings.RUTA_REPORTING)

    def describir(self) -> str:
        return f"SharePoint ReportingFractalia: {mappings.RUTA_REPORTING}"


class FuenteLocal:
    """Contingencia: un .xlsx en disco (misma estructura que la ruta original del .dtsx)."""

    def __init__(self, ruta: Path) -> None:
        self._ruta = ruta

    def leer(self) -> bytes:
        if not self._ruta.exists():
            raise ExtraccionError(f"No existe el archivo local: {self._ruta}")
        return self._ruta.read_bytes()

    def describir(self) -> str:
        return f"archivo local: {self._ruta}"


def valor_ace(valor: Any) -> Any:
    """Valor de una celda tal como lo entrega ACE para una columna DT_WSTR;
    NULL (None) si la celda no es texto."""
    return valor if isinstance(valor, str) else None


def leer_hoja(contenido: bytes, nombre_archivo: str = mappings.ARCHIVO) -> list[Fila]:
    """Lee 'Hoja1' y devuelve una fila por cada fila de datos (bajo la
    cabecera) con las 7 columnas del Excel Source tipadas como ACE."""
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as exc:  # zip corrupto, no es xlsx, etc.
        raise ExtraccionError(f"No se pudo abrir '{nombre_archivo}': {exc}") from exc
    try:
        if mappings.HOJA not in libro.sheetnames:
            raise ExtraccionError(f"'{nombre_archivo}' no tiene la hoja '{mappings.HOJA}' (hojas: {libro.sheetnames}).")
        filas_excel = libro[mappings.HOJA].iter_rows(values_only=True)
        cabecera = next(filas_excel, None) or ()
        indices: dict[str, int] = {}
        for i, nombre in enumerate(cabecera):
            if isinstance(nombre, str) and nombre not in indices:
                indices[nombre] = i  # ante nombres repetidos ACE usa la primera
        faltantes = [c for c in mappings.COLUMNAS_ORIGEN if c not in indices]
        if faltantes:
            raise ExtraccionError(
                f"'{nombre_archivo}' / {mappings.HOJA}: faltan las columnas {faltantes} (cabecera: {cabecera})."
            )
        return [
            {c: valor_ace(fila[indices[c]] if indices[c] < len(fila) else None) for c in mappings.COLUMNAS_ORIGEN}
            for fila in filas_excel
        ]
    finally:
        libro.close()


class Extractor:
    def __init__(self, fuente: FuenteArchivo) -> None:
        self._fuente = fuente

    def extraer(self) -> list[Fila]:
        logger.info("Leyendo %s", self._fuente.describir())
        filas = leer_hoja(self._fuente.leer())
        logger.info("%s fila(s) leidas de '%s'.", len(filas), mappings.HOJA)
        return filas
