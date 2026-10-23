"""Escritura local de resultados en SALIDA_DIR:

    historico_actualizaciones.csv    acumulado de TODAS las corridas (';', UTF-8 con BOM).
                                     Power BI solo conserva un historial limitado por
                                     dataset; este archivo guarda lo que ya no devuelve.
    BI_Service_Actualizaciones.xlsx  foto de la corrida: analisis de cruces (ver
                                     analisis.py) + Datasets, Actualizaciones (historico
                                     completo), Informes, Workspaces y Errores.

Ambos se escriben a un temporal y se reemplazan al final, asi una falla (p.ej.
el Excel abierto) no deja archivos a medias.
"""

from __future__ import annotations

import logging
import os
from pathlib import Path

import pandas as pd
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

from exceptions import SalidaError
from logging_setup import NOMBRE_LOGGER
from transformer import CLAVE_ACTUALIZACION, COLUMNAS_ACTUALIZACIONES, COLUMNAS_FECHA_ACTUALIZACIONES

logger = logging.getLogger(NOMBRE_LOGGER)

ARCHIVO_HISTORICO = "historico_actualizaciones.csv"
ARCHIVO_EXCEL = "BI_Service_Actualizaciones.xlsx"
HOJA_MAPA = "Mapa_cruces"
_SEPARADOR = ";"
_ANCHO_MAX_COLUMNA = 60


def leer_historico(ruta: Path) -> pd.DataFrame:
    if not ruta.exists():
        return pd.DataFrame(columns=COLUMNAS_ACTUALIZACIONES)
    df = pd.read_csv(
        ruta, sep=_SEPARADOR, encoding="utf-8-sig", keep_default_na=False, na_values=[""],
        dtype={c: "string" for c in CLAVE_ACTUALIZACION + ["WORKSPACE_ID"]},
    )
    for col in COLUMNAS_FECHA_ACTUALIZACIONES:
        if col in df:
            df[col] = pd.to_datetime(df[col], errors="coerce", format="ISO8601")
    return df.reindex(columns=COLUMNAS_ACTUALIZACIONES)


def combinar_historico(historico: pd.DataFrame, nuevas: pd.DataFrame) -> pd.DataFrame:
    """Une sin duplicar: si una actualizacion ya existia, gana la version nueva
    (p.ej. paso de 'Unknown' = en curso a 'Completed')."""
    partes = [df for df in (historico, nuevas) if not df.empty]
    if not partes:
        return nuevas.copy()
    df = pd.concat(partes, ignore_index=True)
    clave = df[CLAVE_ACTUALIZACION].astype("string").fillna("")
    df = df[~clave.duplicated(keep="last")]
    return df.sort_values("INICIO_UTC", ascending=False, ignore_index=True)


def _recalcular_locales(df: pd.DataFrame, zona: str) -> pd.DataFrame:
    """*_LOCAL desde *_UTC: si cambia ZONA_HORARIA, las filas antiguas del CSV
    quedan en la misma zona que las nuevas."""
    df = df.copy()
    for utc, local in (("INICIO_UTC", "INICIO_LOCAL"), ("FIN_UTC", "FIN_LOCAL")):
        df[local] = pd.to_datetime(df[utc]).dt.tz_localize("UTC").dt.tz_convert(zona).dt.tz_localize(None)
    return df


def _reemplazar(tmp: Path, destino: Path) -> None:
    try:
        os.replace(tmp, destino)
    except PermissionError as exc:
        tmp.unlink(missing_ok=True)
        raise SalidaError(f"No se pudo reemplazar '{destino}': esta abierto en otro programa (cerrarlo).") from exc


def _formatear_hoja(hoja, df: pd.DataFrame) -> None:
    hoja.freeze_panes = "B2" if hoja.title == HOJA_MAPA else "A2"
    if len(df.columns):
        hoja.auto_filter.ref = hoja.dimensions
    for i, col in enumerate(df.columns, start=1):
        largo = max([len(str(col))] + [len(str(v)) for v in df[col].head(500).tolist()])
        hoja.column_dimensions[get_column_letter(i)].width = min(largo + 2, _ANCHO_MAX_COLUMNA)
    if hoja.title == HOJA_MAPA and len(df):
        rango = f"B2:{get_column_letter(len(df.columns))}{len(df) + 1}"
        hoja.conditional_formatting.add(
            rango, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                  mid_type="num", mid_value=2, mid_color="FFE699",
                                  end_type="max", end_color="F8696B"),
        )


def escribir_excel(ruta: Path, hojas: dict[str, pd.DataFrame]) -> None:
    """Columnas que empiezan con '_' son internas y no se escriben."""
    tmp = ruta.with_name(f"~tmp_{ruta.name}")
    with pd.ExcelWriter(tmp, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss") as writer:
        for nombre, df in hojas.items():
            df = df[[c for c in df.columns if not str(c).startswith("_")]]
            df.to_excel(writer, sheet_name=nombre, index=False)
            _formatear_hoja(writer.sheets[nombre], df)
    _reemplazar(tmp, ruta)


def escribir_csv(ruta: Path, df: pd.DataFrame) -> None:
    tmp = ruta.with_name(f"~tmp_{ruta.name}")
    df.to_csv(tmp, sep=_SEPARADOR, encoding="utf-8-sig", index=False, date_format="%Y-%m-%d %H:%M:%S")
    _reemplazar(tmp, ruta)


def actualizar_historico(nuevas: pd.DataFrame, salida_dir: Path, zona: str) -> tuple[pd.DataFrame, Path, int]:
    """Suma 'nuevas' al CSV historico. Devuelve (historico_completo, ruta_csv, filas_nuevas)."""
    salida_dir.mkdir(parents=True, exist_ok=True)
    ruta_csv = salida_dir / ARCHIVO_HISTORICO
    historico = leer_historico(ruta_csv)
    combinado = _recalcular_locales(combinar_historico(historico, nuevas), zona)
    escribir_csv(ruta_csv, combinado)
    return combinado, ruta_csv, len(combinado) - len(historico)


def escribir_libro(salida_dir: Path, hojas: dict[str, pd.DataFrame]) -> Path:
    salida_dir.mkdir(parents=True, exist_ok=True)
    ruta = salida_dir / ARCHIVO_EXCEL
    escribir_excel(ruta, hojas)
    return ruta
